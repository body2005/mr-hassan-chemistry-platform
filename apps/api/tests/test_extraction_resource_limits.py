"""Pre-allocation budgets and source-fidelity regressions (no runtime fixtures)."""
import io
import math
import zipfile

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from reportlab.pdfgen import canvas

from app.main import app
from app.services.document_parsers import parse_knowledge_file, ocr_image_bytes
from app.services.extraction_limits import (ExtractionLimitError, MAX_RENDER_PIXELS,
    MAX_RENDER_SIDE, check_archive, check_image, render_scale)
from tests.test_extraction_contract import make_teacher_course, login


@pytest.mark.parametrize('size', [(612, 792), (14400, 14400), (12000, 100)])
def test_render_scale_is_bounded_before_bitmap_allocation(size):
    scale = render_scale(*size)
    width, height = [math.ceil(side * scale) for side in size]
    assert max(width, height) <= MAX_RENDER_SIDE
    assert width * height <= MAX_RENDER_PIXELS


@pytest.mark.parametrize('size', [(14401, 100), (float('inf'), 100), (0, 100)])
def test_invalid_page_dimensions_rejected(size):
    with pytest.raises(ExtractionLimitError): render_scale(*size)


def test_pdf_preflight_rejects_oversized_page_before_render(tmp_path):
    target = tmp_path / 'unsafe.pdf'
    page = canvas.Canvas(str(target), pagesize=(15000, 15000))
    page.drawString(10, 10, 'Source content'); page.save()
    with pytest.raises(ExtractionLimitError):
        parse_knowledge_file(filename=target.name, file_path=str(target), mime_type='application/pdf')


def test_pdf_page_count_rejected_before_extraction(tmp_path):
    target = tmp_path / 'many.pdf'; page = canvas.Canvas(str(target))
    for _ in range(101): page.showPage()
    page.save()
    with pytest.raises(ExtractionLimitError):
        parse_knowledge_file(filename=target.name, file_path=str(target), mime_type='application/pdf')


def test_compressed_word_bomb_rejected_before_docx_parser():
    data = io.BytesIO()
    with zipfile.ZipFile(data, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', b'0' * (2 * 1024**2))
    with pytest.raises(ExtractionLimitError): check_archive(data.getvalue())


def test_image_dimensions_checked_without_pixel_decode():
    # Header-only object: no allocation of a 65-megapixel raster.
    class Header: width = 10000; height = 6500; size = (10000, 6500)
    with pytest.raises(ExtractionLimitError): check_image(Header())


def test_ocr_timeout_is_explicit_not_success_with_empty_questions(monkeypatch):
    image = Image.new('RGB', (32, 32), 'white'); data = io.BytesIO(); image.save(data, 'PNG'); image.close()
    calls = []
    def timed_out(*args, **kwargs):
        calls.append(kwargs['timeout']); raise RuntimeError('Tesseract process timeout')
    monkeypatch.setattr('pytesseract.image_to_data', timed_out)
    with pytest.raises(ExtractionLimitError, match='timed out|incomplete'):
        ocr_image_bytes(data.getvalue())
    assert calls == [20]


def test_word_question_units_and_distinct_option_sets_are_preserved(db):
    from docx import Document
    inst, teacher, course = make_teacher_course(db, 'fidelity-audit')
    document = Document()
    stems = ['A body of mass 2 kg accelerates at 3 m/s². Calculate F.',
             'Choose the correct molecular formula.', 'Choose the correct molecular formula.']
    document.add_paragraph('1. ' + stems[0]); document.add_paragraph('A) 6 N'); document.add_paragraph('B) 2 N')
    document.add_paragraph('2. ' + stems[1]); document.add_paragraph('A) H₂O'); document.add_paragraph('B) CO₂')
    document.add_paragraph('3. ' + stems[2]); document.add_paragraph('A) NaCl'); document.add_paragraph('B) NH₃')
    source = io.BytesIO(); document.save(source)
    with TestClient(app) as client:
        login(client, teacher, inst.slug)
        response = client.post('/api/v1/quiz/extract-from-file',
                               headers={'X-CSRF-Token': client.cookies.get('matgar_csrf', '')},
                               data={'course_id': str(course.id)},
                               files={'file': ('fidelity.docx', source.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert response.status_code == 200, response.text
    questions = response.json()['questions']
    assert len(questions) == 3
    assert [q['question_text'] for q in questions] == stems
    assert [[o['text'] for o in q['options']] for q in questions] == [['6 N', '2 N'], ['H₂O', 'CO₂'], ['NaCl', 'NH₃']]
    assert all(q['correct_answer'] is None and q['needs_answer_review'] for q in questions)
    assert response.json()['requires_teacher_approval'] is True
