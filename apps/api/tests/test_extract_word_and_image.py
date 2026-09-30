from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.core.config import get_tesseract_cmd
from app.services.document_parsers import parse_docx_document, parse_image_asset
from app.services import document_parsers
from app.services.exam_text_extractor import segment_exam_document


def test_word_questions_in_tables_preserve_document_order() -> None:
    source = Document()
    source.add_paragraph("س1: ما صيغة الماء؟")
    source.add_paragraph("أ) H2O")
    source.add_paragraph("ب) CO2")
    table = source.add_table(rows=3, cols=1)
    table.cell(0, 0).text = "س2: ما رمز الأكسجين؟"
    table.cell(1, 0).text = "أ) O"
    table.cell(2, 0).text = "ب) H"
    source.add_paragraph("س3: اشرح دور العامل الحفاز.")
    data = BytesIO()
    source.save(data)

    parsed = parse_docx_document(data.getvalue(), "qa-questions.docx")
    text = parsed.pages[0].raw_text
    assert text.index("س1:") < text.index("س2:") < text.index("س3:")
    questions = segment_exam_document(parsed, filename="qa-questions.docx")
    assert len(questions) == 3


def test_standalone_scanned_page_is_marked_for_review() -> None:
    if get_tesseract_cmd() is None:
        pytest.skip("live image OCR requires Tesseract; run this test in the API Docker image")
    fixture = Path(__file__).parent / "fixtures" / "blind_inputs" / "05_biology_first_page.png"
    parsed = parse_image_asset(fixture.read_bytes(), fixture.name)
    assert parsed.extracted_via_ocr
    assert parsed.ocr_pages == [1]
    questions = segment_exam_document(parsed, filename=fixture.name)
    assert questions
    assert all(question["text_source"] == "ocr" for question in questions)
    assert all(question["needs_content_review"] for question in questions)


def test_embedded_word_image_text_is_not_silently_discarded(monkeypatch) -> None:
    fixture = Path(__file__).parent / "fixtures" / "blind_inputs" / "05_biology_first_page.png"
    source = Document()
    source.add_paragraph("نص تمهيدي")
    source.add_picture(str(fixture))
    data = BytesIO()
    source.save(data)
    monkeypatch.setattr(
        document_parsers,
        "ocr_image_bytes",
        lambda _data: ("س1: ما صيغة الماء؟\nأ) H2O\nب) CO2", "synthetic-ocr"),
    )

    parsed = parse_docx_document(data.getvalue(), "embedded-exam.docx")
    assert len(parsed.all_images) == 1
    assert "س1: ما صيغة الماء؟" in parsed.pages[0].raw_text
    assert parsed.pages[0].extracted_via_ocr
    questions = segment_exam_document(parsed, filename="embedded-exam.docx")
    assert len(questions) == 1
    assert questions[0]["needs_content_review"]


def test_word_with_real_scanned_exam_image_flags_ocr_questions() -> None:
    if get_tesseract_cmd() is None:
        pytest.skip("live DOCX image OCR requires Tesseract; run in the API Docker image")
    fixture = Path(__file__).parent / "fixtures" / "blind_inputs" / "05_biology_first_page.png"
    source = Document()
    source.add_picture(str(fixture))
    data = BytesIO()
    source.save(data)

    parsed = parse_docx_document(data.getvalue(), "scanned-exam.docx")
    questions = segment_exam_document(parsed, filename="scanned-exam.docx")
    assert questions
    assert parsed.ocr_pages == [1]
    assert all(question["needs_content_review"] for question in questions)
